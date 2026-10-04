import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

INPUT_CSV = "samples/validation_sample.csv"
OUTPUT_XLSX = "samples/Track Validation Sheet.xlsx"

FONT_NAME = "Arial"

HEADER_FILL = PatternFill(start_color="1DB954", end_color="1DB954", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)

CATEGORY_FILLS = {
    "track": PatternFill(start_color="169C46", end_color="169C46", fill_type="solid"),
    "questions": PatternFill(start_color="4A7FBF", end_color="4A7FBF", fill_type="solid"),
    "features": PatternFill(start_color="6B6B6B", end_color="6B6B6B", fill_type="solid"),
}
CATEGORY_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=12)

BODY_FONT = Font(name=FONT_NAME, size=11)
TRACK_FONT = Font(name=FONT_NAME, size=11, bold=True, color="0563C1", underline="single")
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
SUBTEXT_FONT = Font(name=FONT_NAME, size=10, italic=True, color="555555")

THIN_SIDE = Side(style="thin", color="D9D9D9")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
THICK_SIDE = Side(style="thin", color="A6A6A6")
SECTION_BORDER = Border(left=THICK_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

ROW_FILL_EVEN = PatternFill(start_color="F7FBF8", end_color="F7FBF8", fill_type="solid")
ROW_FILL_ODD = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

PITCH_CLASS_MAP = {
    0: "C", 1: "C♯/D♭", 2: "D", 3: "D♯/E♭", 4: "E", 5: "F",
    6: "F♯/G♭", 7: "G", 8: "G♯/A♭", 9: "A", 10: "A♯/B♭", 11: "B",
}

FEATURE_DEFINITIONS = [
    ("Danceability", "How suitable the track feels for dancing based on rhythm and beat (0 = least danceable, 1 = most danceable)"),
    ("Energy", "How intense or lively the track feels — high energy tracks feel fast, loud, and powerful, while low energy feels calm or mellow (0 to 1)"),
    ("Key", "The musical key the track is written in (e.g., C, D, G)"),
    ("Loudness", "The overall loudness of the track, measured in decibels (dB)"),
    ("Mode", "Whether the track uses a Major or Minor mode"),
    ("Speechiness", "How much spoken word (rather than singing) is present in the track"),
    ("Acousticness", "How likely the track is to have acoustic rather than electronically produced characteristics (0 to 1)"),
    ("Instrumentalness", "How likely it is that the track has no vocals at all"),
    ("Liveness", "How likely it is that the track was recorded in front of a live audience"),
    ("Tempo", "The estimated speed of the track, in beats per minute (BPM)"),
    ("Duration", "The length of the track, in minutes/seconds"),
    ("Valence", "How positive or negative the music sounds emotionally (0 = negative, 1 = positive)"),
    ("Lyric Sentiment", "Overall emotional tone of the lyrics based on their wording (-1 = negative, 0 = neutral, 1 = positive)"),
    ("Alignment Score", "Emotional difference between lyrics and music; values closer to 0 show that the emotional tone of the lyrics and music are similar, while larger differences show greater emotional mismatch."),
    ("Popularity", "How popular the track is on Spotify overall, on a scale from 0 to 100"),
]

TRACK_COLUMNS = [
    ("name", "Name"),
    ("artists", "Artist/s"),
]

RATING_COLUMNS = [
    "Does the track feel danceable to you?",
    "Does the energy level match your impression?",
    "Does the mood (valence) match your impression?",
    "Do the lyrics and music emotionally align to you?",
]

FEATURE_COLUMNS = [
    ("danceability", "Danceability"),
    ("energy", "Energy"),
    ("key", "Key"),
    ("loudness", "Loudness (dB)"),
    ("mode", "Mode"),
    ("speechiness", "Speechiness"),
    ("acousticness", "Acousticness"),
    ("instrumentalness", "Instrumentalness"),
    ("liveness", "Liveness"),
    ("tempo", "Tempo (BPM)"),
    ("duration_ms", "Duration"),
    ("valence", "Valence"),
    ("lyric_sentiment", "Lyric Sentiment"),
    ("alignment_gap", "Alignment Score"),
    ("popularity", "Popularity"),
]

ROW_HEIGHT = 34
DATA_START_ROW = 5

def ms_to_minsec(ms):
    if pd.isna(ms):
        return "—"
    total_sec = int(ms) // 1000
    return f"{total_sec // 60}:{total_sec % 60:02d}"

def key_to_pitch(key):
    if pd.isna(key):
        return "—"
    return PITCH_CLASS_MAP.get(int(key), "—")

def build_definitions_sheet(wb):
    ws = wb.active
    ws.title = "Feature Guide"

    ws["A1"] = "Feature Guide for Listening Validation"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 32

    ws["A2"] = ("Thank you for helping validate these tracks! Below is a plain-language "
                "explanation of each musical characteristic you'll see in the sample sheet.")
    ws["A2"].font = SUBTEXT_FONT
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells("A2:B2")
    ws.row_dimensions[2].height = 35

    ws["A3"] = "Feature"
    ws["B3"] = "Definition"
    for col in ("A3", "B3"):
        ws[col].font = HEADER_FONT
        ws[col].fill = HEADER_FILL
        ws[col].alignment = Alignment(horizontal="center", vertical="center", indent=1)
    ws.row_dimensions[3].height = 22

    for i, (feature, definition) in enumerate(FEATURE_DEFINITIONS, start=4):
        fill = ROW_FILL_EVEN if i % 2 == 0 else ROW_FILL_ODD
        fcell = ws.cell(row=i, column=1, value=feature)
        fcell.font = Font(name=FONT_NAME, bold=True, size=11)
        fcell.fill = fill
        fcell.alignment = Alignment(vertical="center")
        fcell.border = THIN_BORDER

        dcell = ws.cell(row=i, column=2, value=definition)
        dcell.font = BODY_FONT
        dcell.fill = fill
        dcell.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
        dcell.border = THIN_BORDER
        ws.row_dimensions[i].height = 32

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 100
    ws.freeze_panes = "A4"

def build_sample_sheet(wb, df):
    ws = wb.create_sheet("Validation Sample")

    ws["A1"] = "Validation Sample Spotify Tracks"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells("A1:F1")
    ws.row_dimensions[1].height = 32

    ws["A2"] = ("Instructions: Click on the name of a track to watch its lyric video, then select Yes or No for each validation question "
                "by clicking on the down arrow shown on the right side of the selected cell.")
    ws["A2"].font = SUBTEXT_FONT
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = 35

    all_columns = TRACK_COLUMNS + [(None, h) for h in RATING_COLUMNS] + FEATURE_COLUMNS
    total_cols = len(all_columns)

    # ── Row 3: category band (merged spans) ─────────────────────────────
    track_span = len(TRACK_COLUMNS)
    questions_span = len(RATING_COLUMNS)

    bands = [
        ("Spotify Track", 1, track_span, "track"),
        ("Validation Questions", track_span + 1, track_span + questions_span, "questions"),
        ("Audio and Lyrical Features", track_span + questions_span + 1, total_cols, "features"),
    ]
    for label, start_col, end_col, key in bands:
        ws.cell(row=3, column=start_col, value=label)
        if end_col > start_col:
            ws.merge_cells(start_row=3, start_column=start_col, end_row=3, end_column=end_col)
        for c in range(start_col, end_col + 1):
            cell = ws.cell(row=3, column=c)
            cell.fill = CATEGORY_FILLS[key]
            cell.font = CATEGORY_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 26

    # ── Row 4: actual column headers ─────────────────────────────────────
    for col_idx, (_, display) in enumerate(all_columns, start=1):
        cell = ws.cell(row=4, column=col_idx, value=display)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 42

    # ── Dropdown validation for rating columns ("-", "Yes", "No") ──────────
    rating_start_col = track_span + 1
    rating_end_col = track_span + questions_span
    n_rows = len(df)
    dv = DataValidation(type="list", formula1='"-,Yes,No"', allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    if n_rows > 0:
        start_letter = get_column_letter(rating_start_col)
        end_letter = get_column_letter(rating_end_col)
        dv.add(f"{start_letter}{DATA_START_ROW}:{end_letter}{DATA_START_ROW + n_rows - 1}")

    # ── Data rows ─────────────────────────────────────────────────────
    for i, row in enumerate(df.itertuples(index=False)):
        row_idx = DATA_START_ROW + i
        row_dict = row._asdict()
        fill = ROW_FILL_EVEN if row_idx % 2 == 0 else ROW_FILL_ODD
        ws.row_dimensions[row_idx].height = ROW_HEIGHT

        # Track name → hyperlinks to YouTube now
        name = row_dict.get("name")
        youtube_url = row_dict.get("youtube_url")
        name_cell = ws.cell(row=row_idx, column=1)
        name_cell.value = ("▶ " + name) if pd.notna(name) else "—"
        name_cell.border = THIN_BORDER
        name_cell.fill = fill
        name_cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True, indent=1)
        if isinstance(youtube_url, str) and youtube_url.startswith("http"):
            name_cell.hyperlink = youtube_url
            name_cell.font = TRACK_FONT
        else:
            name_cell.font = Font(name=FONT_NAME, bold=True, size=11)

        # Artist(s)
        artists = row_dict.get("artists")
        artist_cell = ws.cell(row=row_idx, column=2)
        artist_cell.value = (
            artists.translate(str.maketrans({'[': '', ']': '', '"': ''}))
            if pd.notna(artists) else "—"
        )
        artist_cell.font = BODY_FONT
        artist_cell.border = THIN_BORDER
        artist_cell.fill = fill
        artist_cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True, indent=1)

        # Rating dropdown cells
        for offset in range(questions_span):
            col_idx = rating_start_col + offset
            cell = ws.cell(row=row_idx, column=col_idx, value="-")
            cell.font = Font(name=FONT_NAME, size=11, bold=True, color="4A7FBF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill = fill
            cell.border = THIN_BORDER

        # Audio/lyrical feature cells
        feature_start_col = rating_end_col + 1
        for offset, (src_col, _) in enumerate(FEATURE_COLUMNS):
            col_idx = feature_start_col + offset
            value = row_dict.get(src_col)
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = SECTION_BORDER if offset == 0 else THIN_BORDER
            cell.fill = fill
            cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
            cell.font = BODY_FONT

            if src_col == "key":
                cell.value = key_to_pitch(value)
            elif src_col == "duration_ms":
                cell.value = ms_to_minsec(value)
            elif src_col == "mode":
                cell.value = "Major" if value == 1 else "Minor" if value == 0 else "—"
            else:
                cell.value = value if pd.notna(value) else "—"

    # ── Column widths ─────────────────────────────────────────────────
    widths = []
    for src_col, display in all_columns:
        if display == "Name":
            widths.append(40)
        elif display == "Artist/s":
            widths.append(30)
        elif display in RATING_COLUMNS:
            widths.append(24)
        elif display in ("Key", "Mode"):
            widths.append(12)
        else:
            sample_vals = [str(v) for v in df.get(src_col, [])] if src_col else []
            widths.append(min(max([len(display)] + [len(v) for v in sample_vals], default=10) + 4, 20))

    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    freeze_col_letter = get_column_letter(rating_end_col + 1)
    ws.freeze_panes = f"{freeze_col_letter}{DATA_START_ROW}"

def main():
    df = pd.read_csv(INPUT_CSV)

    wb = Workbook()
    build_definitions_sheet(wb)
    build_sample_sheet(wb, df)
    wb.save(OUTPUT_XLSX)
    print(f"Saved {OUTPUT_XLSX}")

if __name__ == "__main__":
    main()